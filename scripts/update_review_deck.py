#!/usr/bin/env python
"""Update the Gamma-authored review deck: text only, plus watermark removal.

Deliberately conservative. Every picture, diagram, table geometry, colour and
font in the source deck is left untouched -- the only things this changes are
text runs, the "Made with Gamma" badge, and speaker notes.

Why each edit exists is recorded in EDIT_LOG below, so the diff can be reviewed
as a set of claims rather than a blob.

    python scripts/update_review_deck.py
"""

from __future__ import annotations

import shutil
import zipfile
from copy import deepcopy
from pathlib import Path

from pptx import Presentation

SRC = Path.home() / "Downloads" / "Trust-Gated-Fast-Weight-Updates-for-TTT-E2E-LLMs.pptx"
OUT_DIR = Path(__file__).resolve().parents[1] / "docs" / "review"
OUT = OUT_DIR / "TrustGate_Review_Updated_2026-08-28.pptx"
TRANSCRIPT = OUT_DIR / "TrustGate_Review_Transcript.md"

OLD_FOOTER = "CONFIDENTIAL – INTERNAL RESEARCH – PROVISIONAL PATENT NOT YET FILED"
NEW_FOOTER = (
    "PUBLIC RESEARCH REPOSITORY SINCE 2026-08-01  ·  US FILING DEADLINE ~2027-08-01  "
    "·  SEE DISCLOSURE.md"
)

EDIT_LOG: list[str] = []


def log(what: str) -> None:
    EDIT_LOG.append(what)


# ------------------------------------------------------------------ helpers
def set_lines(text_frame, lines: list[str]) -> None:
    """Rewrite a text frame's paragraphs, keeping each paragraph's formatting.

    The source deck is one run per paragraph throughout, so cloning paragraph 0
    preserves size, weight, colour and the Inter typeface exactly.
    """
    body = text_frame._txBody
    template = deepcopy(text_frame.paragraphs[0]._p)

    while len(text_frame.paragraphs) < len(lines):
        body.append(deepcopy(template))
    while len(text_frame.paragraphs) > len(lines):
        body.remove(text_frame.paragraphs[-1]._p)

    for para, line in zip(text_frame.paragraphs, lines):
        runs = para.runs
        if not runs:
            continue
        runs[0].text = line
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)


def swap(text_frame, old: str, new: str) -> bool:
    changed = False
    for para in text_frame.paragraphs:
        for run in para.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)
                changed = True
    return changed


def shape_by_index(slide, index):
    return list(slide.shapes)[index]


# -------------------------------------------------------------- watermark
def strip_watermark(prs) -> None:
    """Remove the 'Made with Gamma' badge from the shared layout."""
    removed = 0
    for layout in prs.slide_masters[0].slide_layouts:
        for shape in list(layout.shapes):
            if shape.shape_type is not None and "PICTURE" in str(shape.shape_type):
                # The badge is a ~1.17 x 0.28in image pinned to the bottom-right
                # corner of the layout; no real content sits there.
                if shape.top > 6_400_000 and shape.left > 10_500_000:
                    shape._element.getparent().remove(shape._element)
                    removed += 1
    log(f"Removed the Gamma watermark image from the shared layout ({removed} shape).")


def purge_gamma_from_package(path: Path) -> int:
    """Drop the gamma hyperlink relationship and the badge image from the file.

    python-pptx does not garbage-collect relationships when a shape is deleted,
    so the gamma.app URL would otherwise survive inside the saved package even
    though nothing renders it.
    """
    tmp = path.with_suffix(".tmp.pptx")
    dropped = 0

    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(
        tmp, "w", zipfile.ZIP_DEFLATED
    ) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)

            if item.filename == "ppt/media/image-1002-1.png":
                dropped += 1
                continue

            if item.filename.endswith(".rels"):
                text = data.decode("utf-8")
                if "gamma" in text.lower() or "image-1002-1.png" in text:
                    import re

                    before = text
                    text = re.sub(
                        r"<Relationship[^>]*(?:gamma|image-1002-1\.png)[^>]*/>",
                        "",
                        text,
                        flags=re.IGNORECASE,
                    )
                    if text != before:
                        dropped += 1
                    data = text.encode("utf-8")

            zout.writestr(item, data)

    tmp.replace(path)
    return dropped


# ------------------------------------------------------------------ notes
TRANSCRIPT_TEXT: dict[int, str] = {
    1: (
        "Good morning. This project is called TrustGate. The one-line version is that "
        "a new class of language model has started updating its own weights while it is "
        "serving requests, and we think that turns the serving path into a write path "
        "into model state. Our work is the defense for that, and just as importantly, "
        "the experiment that tests whether the threat is real before we build the "
        "defense. I want to flag one thing up front: this repository has been public "
        "since the first of August. That was a deliberate decision, and it means the US "
        "filing window is now a deadline rather than a safety net."
    ),
    2: (
        "Some background on what we are defending. Standard attention costs grow with "
        "the square of the sequence length, which makes very long contexts expensive. "
        "TTT-E2E takes a different route: it treats long context as continual learning. "
        "While the model reads your context, it runs an inner training loop and updates "
        "a small set of what the authors call fast weights. The base weights stay "
        "frozen. The fast weights absorb the context. That is genuinely clever, and it "
        "is why the architecture is interesting. It is also the entire reason we are "
        "standing here."
    ),
    3: (
        "Here is the risk. Because the model learns from the context, the serving path "
        "is now a write path into model state. An attacker does not need to touch your "
        "training data, your base weights, or our defense code. They only need to supply "
        "context, which is the one thing the system is built to accept. And the danger "
        "is cumulative. No single update looks abnormal. The harm arrives as a slow "
        "drift across many ordinary-looking updates, and it shows up later, on unrelated "
        "benign inputs, or as a latent trigger that fires on a phrase the attacker chose."
    ),
    4: (
        "We checked whether anyone had already solved this. They have not. Federated "
        "learning has excellent defenses, but they assume you are aggregating updates "
        "from many clients, and here there is a single stream with nothing to aggregate. "
        "Robust test-time adaptation methods like MedBN defend batch-normalization "
        "statistics in vision models, which is not what a fast-weight LLM update is. "
        "Nothing in the literature does all three of the things we need together: check "
        "every update before it commits, bound how far the weights can drift in total, "
        "and roll back when that bound is breached."
    ),
    5: (
        "This is the survey behind that claim: fifteen recent papers across five "
        "families. I will not read them all. The pattern worth noticing is the Unit "
        "row. Every one of these defends a different unit than we need to. Prompt "
        "injection work defends the prompt. Federated work defends a client gradient. "
        "Offline hardening defends the weights before deployment. Graph poisoning "
        "defends discrete edges. Not one of them defends a continuous fast-weight delta "
        "at the moment it commits, on a single stream, during inference."
    ),
    6: (
        "This table is the same argument in one view. Read the last row against the "
        "others. What is missing across the whole field is a state-integrity controller "
        "sitting at the fast-weight commit point, with a stated guarantee and a recovery "
        "path. That gap is the contribution we are claiming, and it is a systems claim "
        "rather than a detector claim. We are not asserting we can spot poison better "
        "than anyone else. We are asserting that the update can be bounded and undone."
    ),
    7: (
        "This is the full architecture in one picture. Five stages: intercept the "
        "proposed update, score it, decide whether to accept it and charge it against a "
        "drift budget, keep versioned checkpoints so you can roll back, and log "
        "everything for audit. I will walk through each stage, and I will be explicit "
        "each time about which parts are built and tested and which parts are still "
        "design."
    ),
    8: (
        "Stage one, interception. We wrap the exact function where the vendor commits a "
        "fast-weight update, compute the delta, and hand it to the gate. The vendor tree "
        "is a pinned, read-only submodule; everything we do is an overlay plus a "
        "confined, reversible monkeypatch, because that tree ships with no license and "
        "we must not edit it. One finding from this stage is worth your attention. The "
        "vendor throws the adapted weights away. It computes them inside a scan and "
        "never reads the result out. We found that at the desk, before spending anything "
        "on hardware, and it changed the design: we now thread that state ourselves."
    ),
    9: (
        "Stage two, scoring. The primary signal compares the candidate model against a "
        "frozen anchor taken at deployment, on a small rotating probe set, using a "
        "symmetric KL divergence. A cheap gradient-based check acts as a backstop, and a "
        "more expensive influence estimate is reserved for offline auditing. I want to "
        "be direct about status: all three of these are design, not implementation. They "
        "are deliberately not built yet, because if the attack does not work there is no "
        "defense worth building, and we would rather not spend the effort first."
    ),
    10: (
        "Stage three is where the actual guarantee lives. An update is accepted only if "
        "every condition holds at once, including that the accumulated drift plus this "
        "update stays under the budget. Accepted drift adds up, and the breach flag "
        "latches, so once it trips it stays tripped until a window boundary or a "
        "rollback. That latching is the point. The guarantee is about cumulative drift "
        "across a window, not the size of any one update, because the threat is a slow "
        "one. Honest status: the accumulator and the store are built and unit-tested, "
        "but the interceptor does not call them yet. So the bound is designed, not "
        "enforced, and it is blocked by the vendor's fixed scan structure."
    ),
    11: (
        "Stage four, recovery. Trusted states go into a fixed-depth ring buffer, and "
        "rollback is index arithmetic, so it costs the same regardless of how many "
        "updates have happened since. It is constant time, but it is not free: the "
        "buffer sits in GPU memory, and its depth has to cover however long detection "
        "takes. For a billion-parameter model that memory competes directly with "
        "activations. That trade-off is a real design constraint, not a footnote."
    ),
    12: (
        "Stage five, evidence. Every decision is logged as append-only JSON lines so a "
        "run can be reconstructed afterwards. On the evaluation side, the thing I most "
        "want to convey is that we fixed the pass mark before we could see any results. "
        "Effect size at least 0.8, benign degradation at least ten percent, and the "
        "poison stream's perplexity no more than 1.5 times the control's. All three are "
        "required. Those numbers now live in code with a test that reads the frozen "
        "document and fails the build if they ever disagree, so the bar cannot move "
        "quietly."
    ),
    13: (
        "Where we actually are. The scaffold has grown into a complete CPU "
        "implementation: 197 tests, the attack objectives, the stream builders, the "
        "matching discipline, the harness, the report, and continuous integration. What "
        "is still missing is everything that needs the model itself, plus the Phase 2 "
        "gate signals we are deliberately not writing yet. I want to be very clear about "
        "the claim: we have produced no measurement. No GPU has been provisioned, no "
        "checkpoint fetched, and the spend so far is zero against a written cap of three "
        "hundred and twenty-five dollars."
    ),
    14: (
        "Limitations, and I would rather state them than be asked. The core gate signals "
        "are unbuilt. The bounded-drift guarantee is designed but not yet wired in. Our "
        "vendor integration is pinned to one commit and would break if that moved. And "
        "one of our own three pass criteria has a hole we measured: perplexity punishes "
        "scrambled text but rewards repetition, so a poison stream that reuses one "
        "passage repeatedly would slip past the fluency bar. We report span reuse "
        "separately because of that. The next step is not a research question, it is a "
        "purchase: a logging key that the vendor's code path requires before evaluation "
        "will run at all. Everything downstream of it is written and waiting."
    ),
}


# ------------------------------------------------------------------- edits
def apply_text_edits(prs) -> None:
    slides = list(prs.slides)

    # --- footer on every slide -------------------------------------------
    hits = 0
    for slide in slides:
        for shape in slide.shapes:
            if shape.has_text_frame and swap(shape.text_frame, OLD_FOOTER, NEW_FOOTER):
                hits += 1
    log(
        f"Replaced the confidentiality footer on {hits} slides. The old text claimed "
        f"the work was confidential and unfiled; the repository has been deliberately "
        f"public since 2026-08-01, and the same stale wording was already removed from "
        f"pyproject.toml, __init__.py and the README by deliverable D7."
    )

    # --- slide 1 ----------------------------------------------------------
    s1 = slides[0]
    if swap(shape_by_index(s1, 2).text_frame, "TEAM 2", "PROJECT REVIEW"):
        log("Slide 1: replaced the stray label 'TEAM 2' with 'PROJECT REVIEW'.")

    # --- slide 8: the carry finding --------------------------------------
    s8 = slides[7]
    set_lines(
        shape_by_index(s8, 9).text_frame,
        [
            "Vendor TTT-E2E inner loop proposes new state",
            "Update delta: \\Delta\\theta_t = \\theta_{(t+1)}^{prop} - \\theta_t",
            "Interceptor placed at MetaModel.inner_loop_step",
            "Vendor code is read-only, pinned submodule",
            "Overlay plus a confined, reversible monkeypatch",
            "Vendor discards its scan carry — our overlay threads the fast weights (ADR-006)",
        ],
    )
    log(
        "Slide 8: added the vendor scan-carry finding to Stage 1. It is the single most "
        "consequential thing we learned, it belongs in the interception stage, and it "
        "was absent from the deck entirely."
    )

    set_lines(
        shape_by_index(s8, 20).text_frame,
        [
            "The interceptor is non-invasive: vendor code is pinned and read-only. The "
            "vendor also discards its adapted state, so the overlay must carry it forward."
        ],
    )
    log("Slide 8: extended the design principle to state the consequence of the carry gap.")

    # --- slide 10: the guarantee is not yet enforced ----------------------
    s10 = slides[9]
    set_lines(
        shape_by_index(s10, 15).text_frame,
        [
            "The breach flag latches; the budget resets only at a window boundary or after rollback",
            "The guarantee bounds cumulative accepted drift, not merely the norm of one update",
            "STATUS: accumulator and store are built and unit-tested but NOT yet called by "
            "the interceptor — the bound is designed, not enforced (ADR-003)",
        ],
    )
    log(
        "Slide 10: added the enforcement status. The slide presented the drift "
        "mathematics as if it were live; the accumulator and store are built but not "
        "wired into the interceptor, which is the project's most important open caveat."
    )

    # --- slide 12: reference model and threshold pinning ------------------
    s12 = slides[11]
    set_lines(
        shape_by_index(s12, 10).text_frame,
        [
            "Poison versus length-matched benign control",
            "SELECT is the headline; PARAPHRASE and SOFT reported separately",
            "Fluency scored by GPT-2 small — validated by logit equivalence",
            "Difficulty-matched and contiguous arms reported, non-gating",
        ],
    )
    log(
        "Slide 12: named the fluency reference model and its acceptance evidence. The "
        "slide previously said only 'independent reference model', which was true when "
        "written and is now under-specified."
    )

    set_lines(
        shape_by_index(s12, 12).text_frame,
        [
            "Cohen's d \\ge 0.8; relative benign degradation \\ge 10%",
            "Fluency ratio \\le 1.5, gated on the MAX across the five seeds",
            "Thresholds CI-pinned to the frozen document — a bar cannot move silently",
        ],
    )
    log(
        "Slide 12: recorded the max-across-seeds aggregation and the CI threshold pin. "
        "The pre-registration states the fluency bar as one number but the design has "
        "five streams, so the aggregation had to be resolved in writing."
    )

    # --- slide 13: status ------------------------------------------------
    s13 = slides[12]
    set_lines(
        shape_by_index(s13, 0).text_frame,
        ["Current Implementation Status: Phase 1 CPU Implementation Complete"],
    )

    table = None
    for shape in s13.shapes:
        if shape.has_table:
            table = shape.table
            break

    set_lines(table.cell(0, 0).text_frame, ["Implemented and CPU-Tested (197 tests passing)"])
    set_lines(table.cell(0, 1).text_frame, ["Pending: model-gated or deliberately deferred"])
    set_lines(
        table.cell(1, 0).text_frame,
        [
            "Pytree subtraction, addition, selection, global norm",
            "Update interceptor; pass-through and norm-threshold gates",
            "Latching drift accumulator; fixed-depth versioned store",
            "Fast-weight carry across chunks and sequences (ADR-006)",
            "Attack objectives: DEGRADE and TRIGGER, on the victim's own loss",
            "Token corpus and crafted-stream builders (ADR-007)",
            "Poison/control matching discipline — 13 asserted fields",
            "Spike orchestration, verdict report, harness CLI",
            "Independent fluency reference model and scorer",
            "Corruption metrics, clean regression, audit logging",
            "Continuous integration over the CPU suite",
        ],
    )
    set_lines(
        table.cell(1, 1).text_frame,
        [
            "Frozen-anchor consistency gate (Phase 2, by design)",
            "Rotating probe generation (Phase 2, by design)",
            "Uncertainty gate and decision policy (Phase 2, by design)",
            "Offline influence auditor (Phase 2, by design)",
            "MedBN-analogue robust aggregation baseline",
            "Binding the carry overlay to the real vendor model",
            "Adversarial stream search (craft_stream)",
            "TTT-E2E baseline reproduction — needs GPU and checkpoint",
            "Five-seed attack spike and the PROCEED/STOP verdict",
            "Wiring the drift accumulator into the interceptor",
        ],
    )
    log(
        "Slide 13: updated the status table. Test count 31 -> 197; moved the stream "
        "builders, harness, report, CLI, fluency scorer, carry overlay and CI into the "
        "implemented column; kept the Phase 2 gate signals pending and marked them as "
        "deferred by design rather than merely unfinished."
    )

    set_lines(
        shape_by_index(s13, 3).text_frame,
        [
            "Phase 0.4: desk preparation complete — eval command, pre-registered bar, cost model",
            "Phase 1 (CPU): complete — attack machinery, harness and instrument built and tested",
            "Phase 1 (GPU): blocked on a W&B key and checkpoint · Phase 2: only if the spike says PROCEED",
        ],
    )
    log("Slide 13: replaced the timeline, which still described the project as a scaffold.")

    set_lines(
        shape_by_index(s13, 7).text_frame,
        [
            "Runnable end to end on CPU and validated. Before any number: a W&B key (the "
            "vendor authenticates before eval runs), then checkpoint access. Spend $0 of $325."
        ],
    )
    log(
        "Slide 13: replaced the closing sentence with the actual blocker. The deck never "
        "mentioned the W&B key, which is the single thing standing between the project "
        "and its first measurement."
    )

    # --- slide 14: limitations and next steps ----------------------------
    s14 = slides[13]
    set_lines(
        shape_by_index(s14, 5).text_frame,
        [
            "A behavioral gate can produce false rejects or accepts; the uncertainty signal "
            "may miss slow, low-variance poisoning",
            "The threat model assumes the fast-weight update path is exposed to an untrusted stream",
            "The MedBN prior-art comparison still requires a formal line-level review",
            "Perplexity rewards repetition (1.90 vs 19.98 for fluent prose), so the fluency "
            "bar cannot catch a repetitive stream — span reuse is reported separately",
            "Repository public by decision (2026-08-01); absolute-novelty jurisdictions "
            "forfeit, US filing deadline ~2027-08-01",
        ],
    )
    log(
        "Slide 14: removed 'the repository and mechanism are confidential until patent "
        "filing is addressed', which is false, and replaced it with the actual "
        "disclosure posture. Added the measured repetition weakness in our own fluency "
        "criterion."
    )

    set_lines(
        shape_by_index(s14, 8).text_frame,
        [
            "Procure a Weights & Biases entity, project and API key — the one hard "
            "blocker, and a purchase rather than a task",
            "Bind the carry overlay to the real vendor model and implement benign "
            "evaluation",
            "Reproduce the 1B DCLM+Books @8K TTT-E2E baseline against the pre-registered "
            "bracket",
            "Run the pre-registered five-seed SELECT attack spike and record the "
            "PROCEED/STOP verdict",
            "Implement anchor consistency, rotating probes, uncertainty scoring and "
            "policy combination — only if the verdict is PROCEED",
            "Wire the drift accumulator and versioned store into the interceptor",
            "Build and tune the MedBN-analogue robust-gradient baseline",
            "Measure ROC/operating points, clean utility, overhead, memory and rollback",
            "Test adaptive attackers, longer streams, and 3B or larger models",
            "Formalize the bounded-drift proof; complete the prior-art and legal review",
        ],
    )
    log(
        "Slide 14: reordered next steps so the actual blocker leads, and made the "
        "Phase 2 items explicitly conditional on a PROCEED verdict."
    )

    set_lines(
        shape_by_index(s14, 10).text_frame,
        [
            "\"TTT-E2E improves long-context modeling by allowing models to learn while "
            "serving. The same write path creates a security risk. TrustGate addresses "
            "it by validating fast-weight commits, bounding cumulative drift, retaining "
            "trusted versions and enabling rollback. The experiment that decides whether "
            "the threat is real is now built and validated on CPU; the measurement "
            "itself is the next decisive step.\""
        ],
    )

    set_lines(
        shape_by_index(s14, 12).text_frame,
        [
            "\"Phase 1 implementation complete and validated on CPU (197 tests). No "
            "model-backed attack or defense result is claimed; no measurement has been "
            "taken; spend to date $0 of a $325 cap.\""
        ],
    )
    log("Slide 14: updated the conclusion and result-status quotes to the current state.")


def add_notes(prs) -> None:
    for index, slide in enumerate(prs.slides, start=1):
        text = TRANSCRIPT_TEXT.get(index)
        if text:
            slide.notes_slide.notes_text_frame.text = text
    log(f"Added a speaker-notes transcript to all {len(TRANSCRIPT_TEXT)} slides.")


def write_transcript() -> None:
    lines = [
        "# TrustGate project review — presenter transcript",
        "",
        "> Generated by `scripts/update_review_deck.py` alongside",
        "> `TrustGate_Review_Updated_2026-08-28.pptx`. The same text is embedded in each",
        "> slide's speaker notes, so the deck and this file cannot drift apart.",
        "",
        "Roughly 11–13 minutes at a normal speaking pace.",
        "",
        "---",
        "",
    ]
    titles = {
        1: "Title",
        2: "Background — why TTT-E2E exists",
        3: "The security risk — the serving path is a write path",
        4: "Closing the gap — why existing defenses fall short",
        5: "Literature survey — 15 papers, five families",
        6: "Comparative analysis — where TrustGate fits",
        7: "Architecture pipeline",
        8: "Stage 1 — capturing the fast-weight update",
        9: "Stage 2 — multi-signal validation",
        10: "Stage 3 — acceptance policy and drift control",
        11: "Stage 4 — O(1) rollback",
        12: "Stage 5 — auditability and evaluation",
        13: "Current implementation status",
        14: "Limitations, future work, conclusion",
    }
    for index in sorted(TRANSCRIPT_TEXT):
        lines += [f"## Slide {index} — {titles[index]}", "", TRANSCRIPT_TEXT[index], "", "---", ""]

    lines += [
        "## Questions you should expect",
        "",
        "**“Do you have any results?”**  No, and we say so on two slides. No GPU has been",
        "provisioned and no measurement taken. That is a deliberate ordering: the pass",
        "mark was fixed before the experiment could produce a number, and the machinery",
        "was validated before it was trusted.",
        "",
        "**“Why is a logging key blocking a machine-learning experiment?”**  The vendor",
        "authenticates and queries its logging service before the evaluation branch is",
        "reached, and the result drives a resume decision. Disabling logging does not",
        "avoid that path, and working around it would mean editing a vendored tree that",
        "ships with no license.",
        "",
        "**“Is the bounded-drift guarantee working?”**  Not yet. The accumulator and the",
        "versioned store are implemented and unit-tested, but the interceptor does not",
        "call them, so the bound is designed rather than enforced. It is blocked by the",
        "vendor's fixed scan structure, and it is recorded as a dated correction rather",
        "than quietly carried as done.",
        "",
        "**“What if the attack does not work?”**  Then the project stops, and that is a",
        "complete result. The criterion was pre-registered precisely so that outcome",
        "cannot be rationalised away afterwards. A STOP verdict would be scoped to the",
        "configuration that produced it — 8K context and eight inner steps — not to",
        "TTT-E2E in general.",
        "",
    ]
    TRANSCRIPT.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"source deck not found: {SRC}")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC, OUT)

    prs = Presentation(OUT)
    strip_watermark(prs)
    apply_text_edits(prs)
    add_notes(prs)
    prs.save(OUT)

    dropped = purge_gamma_from_package(OUT)
    log(f"Purged the gamma.app relationship and badge image from the package ({dropped} entries).")

    write_transcript()

    print(f"wrote {OUT}")
    print(f"wrote {TRANSCRIPT}")
    print("\nEdits applied:")
    for i, entry in enumerate(EDIT_LOG, 1):
        print(f"  {i:2d}. {entry}")


if __name__ == "__main__":
    main()
